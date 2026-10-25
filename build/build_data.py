#!/usr/bin/env python3
"""
Gharpital data build.

Inputs (CSV, editable on github.com unless noted):
  data/raw/nmc_doctors.csv          NMC register, written by the scraper (do not hand-edit)
  data/sources/osm_places.csv       places from OpenStreetMap (written by the import; do not hand-edit)
  data/sources/osm_localities.csv   towns/villages from OpenStreetMap, for placing doctors
  data/hospitals.csv                your own places, or fixes to OSM ones (same hospital_id wins)
  data/doctor_hospital.csv          which doctor works at which place (optional)
  data/corrections.csv              hand fixes to doctors; always applied last

Outputs (read by the app; regenerated every build, never hand-edit):
  data/build/meta.json                    counts, provinces/districts with map boxes, categories
  data/build/places/<Province>.json       places in that province (compact rows)
  data/build/doctors/<District>.json      doctors whose registered address is in that district
  data/build/doctors-index.json           every doctor, compact, for name/NMC search and the editor
  data/build/hospitals.json               every place as an object (editor + place pages)
  data/export/gharpital.xlsx              Excel copy for browsing

Run:  python build/build_data.py
"""
import csv
import difflib
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "build"))
from nepal_geo import PROVINCES, DISTRICT_TO_PROVINCE, match_district  # noqa: E402

DATA = ROOT / "data"
OUT = DATA / "build"
EXPORT = DATA / "export"

CATEGORIES = ["hospital", "clinic", "dentist", "optometry", "specialized", "pharmacy", "ayurveda"]
SPECIALIZED = ["Cancer", "Kidney", "Heart", "Orthopaedic", "Children's", "Maternity"]
LEVELS = ["Major", "Provincial", "District", "Primary"]

# Doctor specialty categories (from NMC speciality / degree). First match wins.
SPECIALTY_RULES = [
    ("Cardiology", r"cardi"), ("Neurosurgery", r"neuro.?surg"), ("Neurology", r"neurolog"),
    ("Orthopedics", r"ortho"), ("Obstetrics & Gynecology", r"gyn|obstet|\bobg|ob/?gyn"),
    ("Pediatrics", r"p(a)?ediatric|child health"), ("Dermatology", r"dermat|venereo|skin"),
    ("Ophthalmology", r"ophthal|\beye\b"), ("ENT", r"\bent\b|otorhino|oto-rhino|head and neck"),
    ("Psychiatry", r"psychiat"), ("Radiology", r"radio(logy|diag)|imaging"), ("Anesthesiology", r"an(a)?esth"),
    ("Pathology", r"patholog"), ("Oncology", r"oncolog|cancer"), ("Nephrology", r"nephro"), ("Urology", r"urolog"),
    ("Gastroenterology", r"gastro|hepato"), ("Pulmonology", r"pulmon|chest|respirat"),
    ("Endocrinology", r"endocrin|diabet"), ("General Surgery", r"surg"),
    ("Internal Medicine", r"internal medicine|general medicine|\bmd\b.*medicine"),
    ("Emergency Medicine", r"emergency"), ("Family Medicine", r"family|general practice|\bgp\b|mdgp"),
    ("Community Medicine", r"community|public health|preventive"), ("Microbiology", r"microbio"),
    ("Biochemistry", r"biochem"), ("Pharmacology", r"pharmacol"), ("Forensic Medicine", r"forensic"),
    ("Physiology", r"physiolog"), ("Anatomy", r"anatom"),
    ("Dentistry", r"dent|\bbds\b|\bmds\b|oral|maxillofac|orthodon|periodon|prosthodon|endodon"),
]


def read_csv(path):
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8-sig") as f:
        return [{(k or "").strip(): (v or "").strip() for k, v in row.items()} for row in csv.DictReader(f)]


def categorize(speciality, degree):
    if speciality:
        for cat, pat in SPECIALTY_RULES:
            if re.search(pat, speciality.lower()):
                return cat
    text = f"{speciality} {degree}".lower()
    for cat, pat in SPECIALTY_RULES:
        if re.search(pat, text):
            return cat
    if re.search(r"mbbs|equivalent to mbbs|\bmd\b", text):
        return "General Physician"
    return "Not listed"


def norm_place(s):
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z]", "", s)


# ---------------- geography helpers ----------------
def district_shapes():
    gj = json.loads((DATA / "geo" / "districts.geojson").read_text())
    out = {}
    for f in gj["features"]:
        g = f["geometry"]
        polys = [g["coordinates"]] if g["type"] == "Polygon" else g["coordinates"]
        pts = [pt for poly in polys for pt in poly[0]]
        xs, ys = [p[0] for p in pts], [p[1] for p in pts]
        # area-weighted centroid of the outer rings
        a = cx = cy = 0.0
        for poly in polys:
            ring = poly[0]
            for i in range(len(ring) - 1):
                x0, y0 = ring[i][:2]
                x1, y1 = ring[i + 1][:2]
                c = x0 * y1 - x1 * y0
                a += c
                cx += (x0 + x1) * c
                cy += (y0 + y1) * c
        centroid = (cy / (3 * a), cx / (3 * a)) if a else ((min(ys) + max(ys)) / 2, (min(xs) + max(xs)) / 2)
        out[f["properties"]["district"]] = {
            "bbox": [round(min(ys), 4), round(min(xs), 4), round(max(ys), 4), round(max(xs), 4)],
            "centroid": [round(centroid[0], 4), round(centroid[1], 4)],
        }
    return out


class Gazetteer:
    """Finds a registered locality (e.g. 'Khaireni Tar', district 'Tanahun') on the map."""
    RANK = {"city": 0, "town": 1, "admin": 2, "suburb": 3, "village": 4, "quarter": 5, "neighbourhood": 6,
            "hamlet": 7, "locality": 8}

    def __init__(self, rows):
        self.by_district = defaultdict(dict)  # district -> norm name -> (rank, lat, lng)
        self.national = defaultdict(list)    # norm name -> [(district, rank, lat, lng)]
        for r in rows:
            try:
                lat, lng = float(r["lat"]), float(r["lng"])
            except (KeyError, ValueError):
                continue
            rank = self.RANK.get(r.get("kind", ""), 9)
            for nm in {r.get("name", ""), r.get("name_en", "")}:
                k = norm_place(nm)
                if len(k) < 3:
                    continue
                cur = self.by_district[r["district"]].get(k)
                if cur is None or rank < cur[0]:
                    self.by_district[r["district"]][k] = (rank, lat, lng)
                self.national[k].append((r["district"], rank, lat, lng))

    def find(self, place, district):
        k = norm_place(place)
        if len(k) < 3:
            return None
        if district:
            names = self.by_district.get(district, {})
            if k in names:
                return names[k][1:], district
            close = difflib.get_close_matches(k, names.keys(), n=1, cutoff=0.86)
            if close:
                return names[close[0]][1:], district
            return None
        hits = self.national.get(k, [])
        dists = {h[0] for h in hits}
        if len(dists) == 1:  # unambiguous across Nepal
            best = min(hits, key=lambda h: h[1])
            return (best[2], best[3]), best[0]
        return None


# ---------------- main ----------------
def main():
    warnings = []
    shapes = district_shapes()
    raw = read_csv(DATA / "raw" / "nmc_doctors.csv")
    gaz = Gazetteer(read_csv(DATA / "sources" / "osm_localities.csv"))

    # ---- places: OSM import, older hospital-only import, then your hospitals.csv on top ----
    merged = {}
    osm = read_csv(DATA / "sources" / "osm_places.csv")
    if not osm:  # before the first places import, fall back to the hospital-only file
        osm = [{**h, "category": "hospital"} for h in read_csv(DATA / "sources" / "osm_hospitals.csv")]
    for h in osm:
        if h.get("hospital_id"):
            merged[h["hospital_id"]] = h
    for h in read_csv(DATA / "hospitals.csv"):
        if h.get("hospital_id"):
            merged[h["hospital_id"]] = {**merged.get(h["hospital_id"], {}), **{k: v for k, v in h.items() if v}}

    places = {}
    by_name = {}
    for h in merged.values():
        hid = h["hospital_id"]
        if h.get("type", "").lower() == "hide" or h.get("category", "").lower() == "hide":
            continue
        district = match_district(h.get("district", "")) or h.get("district", "")
        province = DISTRICT_TO_PROVINCE.get(district, h.get("province", ""))
        try:
            lat, lng = float(h["lat"]), float(h["lng"])
        except (KeyError, ValueError):
            warnings.append(f"place {hid} ({h.get('name')}) has no valid location; skipped")
            continue
        cat = (h.get("category") or "hospital").lower()
        if cat not in CATEGORIES:
            warnings.append(f"place {hid}: unknown category '{cat}', treated as hospital")
            cat = "hospital"
        places[hid] = {
            "id": hid, "name": h.get("name", ""), "name_ne": h.get("name_ne", ""), "category": cat,
            "level": h.get("level", "") if cat == "hospital" or h.get("level") == "Health post" else "",
            "specialty": h.get("specialty", "") if cat == "specialized" else "",
            "ownership": h.get("ownership", ""), "province": province, "district": district,
            "address": h.get("address", ""), "lat": round(lat, 6), "lng": round(lng, 6),
            "phone": h.get("phone", ""), "website": h.get("website", ""),
            "hours": h.get("opening_hours", ""), "emergency": h.get("emergency", ""),
        }
        by_name[h.get("name", "").lower()] = hid

    def resolve_place(value):
        return value if value in places else by_name.get((value or "").lower())

    # ---- doctors ----
    doctors = {}
    for r in raw:
        key = r.get("nmc_no", "")
        if not key:
            continue
        doctors[key] = {
            "nmc_no": key, "name": re.sub(r"\s+", " ", r.get("full_name", "")).strip(),
            "gender": r.get("gender", ""), "degree": r.get("degree", ""),
            "speciality_raw": r.get("speciality", ""), "specialty": "", "photo_url": "", "bio": "",
            "reg_place": r.get("reg_place", ""), "reg_district": r.get("reg_district", ""),
            # NMC marks deregistered doctors (death, cancellation...) in Remarks; keep them out of the app
            "hidden": bool(re.search(r"dereg|cancel|struck|suspend|revok|death|deceased", r.get("remarks", ""), re.I)),
        }

    links = defaultdict(dict)
    for l in read_csv(DATA / "doctor_hospital.csv"):
        n, hid = l.get("nmc_no"), resolve_place(l.get("hospital_id", ""))
        if n and hid:
            links[n][hid] = {"id": hid, "department": l.get("department", ""),
                             "opd_days": l.get("opd_days", ""), "opd_time": l.get("opd_time", "")}

    last_place = {}
    for c in read_csv(DATA / "corrections.csv"):
        n, field, val = c.get("nmc_no", ""), c.get("field", "").lower(), c.get("new_value", "")
        if not n or not field:
            continue
        if n not in doctors:
            if n.upper().startswith("NEW") or field == "name":
                doctors[n] = {"nmc_no": n, "name": "", "gender": "", "degree": "", "speciality_raw": "",
                              "specialty": "", "photo_url": "", "bio": "", "reg_place": "", "reg_district": "",
                              "hidden": False}
            else:
                warnings.append(f"corrections.csv: NMC {n} not in register (field {field})")
                continue
        d = doctors[n]
        if field == "hide":
            d["hidden"] = val.lower() in ("yes", "y", "true", "1")
        elif field == "hospital":
            hid = resolve_place(val)
            if hid:
                links[n].setdefault(hid, {"id": hid, "department": "", "opd_days": "", "opd_time": ""})
                last_place[n] = hid
            else:
                warnings.append(f"corrections.csv: place '{val}' not found")
        elif field == "remove_hospital":
            links[n].pop(resolve_place(val), None)
        elif field in ("department", "opd_days", "opd_time"):
            target = last_place.get(n)
            for hid, link in links[n].items():
                if target is None or hid == target:
                    link[field] = val
        elif field in ("district", "reg_district"):
            d["reg_district"] = match_district(val) or val
        elif field in ("place", "reg_place"):
            d["reg_place"] = val
        elif field in ("name", "gender", "degree", "specialty", "photo_url", "bio"):
            d[field] = val
        elif field == "speciality":
            d["specialty"] = val
        else:
            warnings.append(f"corrections.csv: unknown field '{field}'")

    # ---- locate doctors by registered address (rounded to ~1 km) ----
    precision = Counter()
    for d in doctors.values():
        if not d["specialty"]:
            d["specialty"] = categorize(d["speciality_raw"], d["degree"])
        d["hospitals"] = list(links.get(d["nmc_no"], {}).values())
        d["loc"] = None
        district = match_district(d["reg_district"]) or ""
        hit = gaz.find(d["reg_place"], district) if d["reg_place"] else None
        if hit:
            (lat, lng), district = hit
            d["loc"] = (round(lat, 2), round(lng, 2), "p")
            precision["locality"] += 1
        elif district in shapes:
            lat, lng = shapes[district]["centroid"]
            d["loc"] = (round(lat, 2), round(lng, 2), "d")
            precision["district"] += 1
        else:
            precision["unknown"] += 1
        d["reg_district"] = district

    visible = [d for d in doctors.values() if not d["hidden"] and d["name"]]
    visible.sort(key=lambda d: (int(re.sub(r"\D", "", d["nmc_no"]) or 10**9), d["nmc_no"]))

    # ---- write ----
    OUT.mkdir(parents=True, exist_ok=True)
    for sub in ("places", "doctors", "province"):
        folder = OUT / sub
        if folder.exists():
            for old in folder.glob("*.json"):
                old.unlink()
            if sub == "province":
                folder.rmdir()
    (OUT / "places").mkdir(exist_ok=True)
    (OUT / "doctors").mkdir(exist_ok=True)

    # places per province: [id, name, name_ne, category, level, specialty, lat, lng, district, phone,
    #                       website, hours, emergency, ownership, address]
    per_prov = defaultdict(list)
    for p in places.values():
        if p["province"] in PROVINCES:
            per_prov[p["province"]].append([p["id"], p["name"], p["name_ne"], p["category"], p["level"],
                                            p["specialty"], p["lat"], p["lng"], p["district"], p["phone"],
                                            p["website"], p["hours"], p["emergency"], p["ownership"], p["address"]])
    for prov in PROVINCES:
        write_json(OUT / "places" / f"{prov}.json", per_prov.get(prov, []))

    # doctors per registered district: [nmc, name, degree, specialty, lat, lng, precision]
    per_dist = defaultdict(list)
    for d in visible:
        if d["loc"]:
            lat, lng, prec = d["loc"]
            per_dist[d["reg_district"]].append([d["nmc_no"], d["name"], d["degree"], d["specialty"], lat, lng, prec])
    for dist, rows in per_dist.items():
        write_json(OUT / "doctors" / f"{dist}.json", rows)

    # search index: [nmc, name, gender, degree, specialty, photo, [place ids], registered district]
    write_json(OUT / "doctors-index.json",
               [[d["nmc_no"], d["name"], d["gender"][:1], d["degree"], d["specialty"], d["photo_url"],
                 [l["id"] for l in d["hospitals"]], d["reg_district"]] for d in visible])
    write_json(OUT / "hospitals.json", sorted(places.values(), key=lambda p: (p["province"], p["district"], p["name"])))

    cat_counts = Counter(p["category"] for p in places.values())
    meta = {
        "built_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "totals": {
            "doctors": len(visible),
            "doctors_located": sum(1 for d in visible if d["loc"]),
            "places": len(places),
            **{f"places_{c}": cat_counts.get(c, 0) for c in CATEGORIES},
        },
        "provinces": [
            {"name": p,
             "bbox": merge_bbox([shapes[d]["bbox"] for d in ds if d in shapes]),
             "districts": [{"name": d, "bbox": shapes[d]["bbox"] if d in shapes else None,
                            "doctors": len(per_dist.get(d, []))} for d in ds]}
            for p, ds in PROVINCES.items()
        ],
        "specialized": SPECIALIZED,
    }
    write_json(OUT / "meta.json", meta, pretty=True)
    export_excel(visible, places)

    print(f"Built: {len(visible)} doctors ({meta['totals']['doctors_located']} placed by registered address: "
          f"{precision['locality']} by locality, {precision['district']} by district only), "
          f"{len(places)} places {dict(cat_counts)}")
    for w in warnings[:40]:
        print("WARNING:", w)
    if len(warnings) > 40:
        print(f"... and {len(warnings) - 40} more warnings")


def merge_bbox(boxes):
    if not boxes:
        return None
    return [min(b[0] for b in boxes), min(b[1] for b in boxes), max(b[2] for b in boxes), max(b[3] for b in boxes)]


def write_json(path, obj, pretty=False):
    with path.open("w", encoding="utf-8") as f:
        if pretty:
            json.dump(obj, f, ensure_ascii=False, indent=2)
        else:
            json.dump(obj, f, ensure_ascii=False, separators=(",", ":"))


def export_excel(doctors, places):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        print("openpyxl not installed; skipping Excel export")
        return
    EXPORT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2A4B58")

    def sheet(ws, headers, rows, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Doctors"
    sheet(ws, ["NMC No", "Name", "Gender", "Degree", "Specialty", "NMC speciality (raw)", "Registered district"],
          [[d["nmc_no"], d["name"], d["gender"], d["degree"], d["specialty"], d["speciality_raw"], d["reg_district"]]
           for d in doctors], [10, 32, 9, 30, 22, 26, 18])
    sheet(wb.create_sheet("Places"),
          ["ID", "Name", "Category", "Level", "Specialty", "Province", "District", "Lat", "Lng", "Phone", "Website"],
          [[p["id"], p["name"], p["category"], p["level"], p["specialty"], p["province"], p["district"],
            p["lat"], p["lng"], p["phone"], p["website"]] for p in places.values()],
          [18, 36, 12, 12, 12, 14, 16, 10, 10, 16, 30])
    wb.save(EXPORT / "gharpital.xlsx")


if __name__ == "__main__":
    sys.exit(main())
